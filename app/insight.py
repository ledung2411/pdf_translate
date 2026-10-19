"""Tóm tắt bài báo theo các câu hỏi nghiên cứu quan trọng.

Hai chế độ, đều chạy trên máy:

- Có LLM cục bộ (Ollama / LM Studio / llama.cpp): với mỗi câu hỏi, chọn các câu liên quan
  nhất trong bài (theo mục và từ khóa) rồi nhờ LLM trả lời bằng tiếng Việt, chỉ dựa trên
  trích đoạn — ngữ cảnh mỗi lần gọi nhỏ nên chạy được với mô hình 7B, cửa sổ ngữ cảnh ngắn.
- Không có LLM: trích các câu bằng chứng cho từng câu hỏi (dịch sang tiếng Việt bằng mô hình
  dịch đang dùng), kèm dữ kiện tự nhận ra: link mã nguồn, tên dataset, các phương pháp so sánh.
"""

from __future__ import annotations

import html
import logging
import re
from collections import Counter
from dataclasses import dataclass, field
from typing import Callable

from .ir import Document, Region
from .pipeline import MATH_END, MATH_START, SUP_END, SUP_START
from .scientific import REFERENCES_RE, split_sentences, translate_heading

log = logging.getLogger(__name__)
StatusFn = Callable[[str], None]

# ---------------------------------------------------------------------------
# Chia bài thành các mục
# ---------------------------------------------------------------------------

# Thứ tự quan trọng: "Comparison to State-of-the-Art Methods" là thí nghiệm, không phải phương pháp.
_SECTION_KINDS = [
    ("abstract", r"abstract|summary"),
    # "Why Self-Attention", "Rationale", "Design choices": lập luận vì sao chọn giải pháp
    ("rationale", r"\bwhy\b|rationale|justification|intuition|motivation for|design choices?"),
    ("intro", r"introduction|background|motivation|overview"),
    ("related", r"related work|prior work|literature|previous work|preliminar"),
    ("conclusion", r"conclusion|concluding|summary and"),
    ("discussion", r"discussion|limitation|future"),
    ("data", r"dataset|data|material|participant|subject|cohort|acquisition|preprocessing"),
    ("experiment", r"experiment|evaluation|implementation|setup|setting|benchmark|ablation|result|comparison|"
                   r"analys|assessment|validation|efficiency|generali[sz]ation|test|application"),
    ("method", r"method|approach|model|architecture|framework|proposed|design|algorithm|network|formulation|loss|input"),
]
_STOP_SECTIONS = re.compile(r"^\s*(\d+(\.\d+)*\.?\s*)?(acknowledg|references|bibliography|funding|conflict of interest|"
                            r"author contributions|data availability statement|ethics)", re.I)
_NUM_HEADING = re.compile(r"^\s*(\d+(?:\.\d+)*|[IVX]+|[A-Z])[.\s|]+\s*([A-Z].{2,90})$")


def _section_kind(title: str) -> str:
    t = re.sub(r"^\s*(\d+(\.\d+)*\.?|[IVX]+\.)\s*", "", title).lower()
    for kind, pat in _SECTION_KINDS:
        if re.search(pat, t):
            return kind
    return "other"


@dataclass
class Section:
    title: str  # tiêu đề tiếng Anh
    kind: str
    regions: list[Region] = field(default_factory=list)


def _is_heading(r: Region, body_size: float) -> bool:
    text = r.plain.strip()
    if not text or len(text) > 110 or len(r.line_rects) > 3 or text.endswith((".", ",", ";")):
        return False
    if translate_heading(text) is not None:
        return True
    if _NUM_HEADING.match(text) and (r.bold or r.font_size >= body_size * 1.05):
        return True
    return r.bold and r.font_size >= body_size * 1.1 and len(text.split()) <= 12


_FOOTNOTE_MARK = re.compile(r"^\s*([∗*†‡§¶‖]|\d{1,2}\s*[A-Z]|\d{1,2}\s*https?://)")


def _is_footnote(r: Region, body_size: float, page_height: float) -> bool:
    """Chú thích chân trang (đóng góp tác giả, nơi công tác, tên hội nghị): chữ nhỏ ở cuối trang
    hoặc mở đầu bằng ký hiệu chú thích. Chú thích hình/bảng không tính."""
    text = r.plain.strip()
    if r.kind != "text" or not text or _CAPTION.match(text) or r.font_size >= body_size * 0.9:
        return False
    return r.bbox.y0 > page_height * 0.7 or bool(_FOOTNOTE_MARK.match(text))


def split_sections(ir: Document) -> list[Section]:
    """Các mục của bài theo thứ tự đọc (dừng ở Tài liệu tham khảo / Lời cảm ơn)."""
    from .export_docx import reading_order
    import pymupdf

    texts = [r for r in ir.regions if r.kind == "text"]
    sizes = sorted(r.font_size for r in texts if len(r.plain) > 200) or [10.0]
    body = sizes[len(sizes) // 2]
    sections = [Section("(Đầu bài)", "front")]
    stopped = False
    for p in ir.pages:
        items = [(pymupdf.Rect(r.bbox), r) for r in p.regions if r.kind in ("text", "math", "table")]
        for r in reading_order(items, p.width):
            if stopped:
                break
            text = r.plain.strip()
            if r.kind == "text" and (_STOP_SECTIONS.match(text) or REFERENCES_RE.match(text)) and len(text) < 60:
                stopped = True
                break
            if _is_footnote(r, body, p.height):
                continue
            if r.kind == "text" and _is_heading(r, body):
                kind = _section_kind(text)
                # Trước Tóm tắt / Giới thiệu: tên bài, tên tạp chí... là phần đầu bài.
                if kind not in ("abstract", "intro") and not any(x.kind in ("abstract", "intro") for x in sections):
                    kind = "front"
                sections.append(Section(text, kind))
                continue
            if r.kind == "text" and re.match(r"^(ABSTRACT|Abstract)\b", text) \
                    and not any(x.kind in ("intro", "method", "abstract") for x in sections):
                sections.append(Section("Abstract", "abstract"))
            sections[-1].regions.append(r)
        if stopped:
            break
    # Phần đầu bài có "Abstract" ở giữa đoạn → coi là tóm tắt
    if sections and sections[0].kind == "front":
        for r in sections[0].regions:
            if re.search(r"\babstract\b", r.plain[:40], re.I):
                sections[0].kind = "abstract"
    return [s for s in sections if s.regions]


def with_latex(r: Region, text: str | None = None) -> str:
    """Văn bản của vùng, công thức nội dòng ghi bằng LaTeX ``$...$``."""
    text = r.text if text is None else text

    def rep(m: re.Match) -> str:
        raw = m.group(1).strip()
        tex = r.math_latex.get(raw)
        return f" ${tex}$ " if tex else f" {raw} "

    text = re.sub(f"{MATH_START}(.*?){MATH_END}", rep, text, flags=re.S)
    text = text.replace(SUP_START, "^").replace(SUP_END, "")
    return re.sub(r"[ \t]+", " ", text).strip()


# ---------------------------------------------------------------------------
# Bộ câu hỏi nghiên cứu
# ---------------------------------------------------------------------------


@dataclass
class Question:
    key: str
    title: str  # tiêu đề tiếng Việt
    ask: str  # câu hỏi gửi LLM
    pattern: str  # từ khóa (regex, tiếng Anh) để chọn câu bằng chứng
    sections: dict[str, float]  # mục ưu tiên → hệ số
    k: int = 4  # số câu bằng chứng (chế độ không LLM)


QUESTIONS: list[Question] = [
    Question("problem", "Bài toán & tình trạng hiện tại",
             "Bài báo giải quyết bài toán gì? Các phương pháp hiện có đang làm được gì và còn vướng mắc ở đâu?",
             r"\b(however|challenge|challenging|limitation|remain|existing|current|conventional|problem|difficult|"
             r"lack|suffer|bottleneck|struggle|still|unsolved|open question|gap)\b",
             {"abstract": 1.5, "intro": 2.0, "related": 1.2}),
    Question("urgency", "Tính cấp thiết / động lực",
             "Vì sao vấn đề này quan trọng và cần giải quyết ngay (ý nghĩa thực tiễn, nhu cầu, hậu quả nếu không giải quyết)?",
             r"\b(important|importance|critical|crucial|essential|urgent|need|needed|demand|growing|increasingly|"
             r"widely|vital|key role|significant|clinical|real-world|practical|cost|time-consuming|expensive|motivat)\w*",
             {"abstract": 1.5, "intro": 2.0}),
    Question("contribution", "Đóng góp chính",
             "Liệt kê các đóng góp chính / điểm mới của bài báo.",
             r"\b(we propose|we present|we introduce|we develop|we design|our contributions?|contributions? (of|are)|"
             r"novel|for the first time|first to|this (paper|work|study)|in this (paper|work|study)|main contribution)",
             {"abstract": 2.0, "intro": 1.8, "conclusion": 1.2}),
    Question("origin", "Nguồn gốc ý tưởng",
             "Ý tưởng của phương pháp bắt nguồn / kế thừa từ đâu (công trình, kỹ thuật, quan sát nào)?",
             r"\b(inspired|motivated by|based on|building on|builds on|extending|extends|extend the|borrow|adopt|"
             r"adapt|derived from|similar to|in the spirit of|draws? on|following (the|prior|previous|recent))\w*",
             {"intro": 1.5, "related": 2.0, "method": 1.3, "experiment": 0.2}),
    Question("architecture", "Phương pháp / kiến trúc",
             "Mô tả phương pháp hoặc kiến trúc đề xuất: gồm những thành phần/khối nào, hoạt động ra sao, huấn luyện thế nào?",
             r"\b(architecture|module|network|framework|encoder|decoder|layer|block|pipeline|stage|branch|stream|"
             r"backbone|attention|loss|train|trained|optimi[sz]|consist|composed|comprise)\w*",
             {"method": 2.2, "abstract": 1.2}, k=5),
    Question("io", "Đầu vào – đầu ra",
             "Mô hình/phương pháp nhận đầu vào là gì (dạng dữ liệu, kích thước) và cho ra đầu ra gì?",
             r"\b(input|inputs|output|outputs|takes as|take as|produces?|predicts?|generates?|given an?|maps?|"
             r"returns?|yields?|outputs?)\b",
             {"method": 2.0, "abstract": 1.3, "data": 1.2}),
    Question("why", "Vì sao giải pháp phù hợp",
             "Vì sao giải pháp đề xuất phù hợp với bài toán (lập luận, tính chất của dữ liệu/bài toán mà phương pháp khai thác)?",
             r"\b(because|since|therefore|thus|hence|enables?|allows?|ensures?|captur|exploit|leverag|suitable|"
             r"naturally|advantage|benefit|robust|so that|in order to|designed to)\w*",
             {"rationale": 2.6, "method": 1.8, "intro": 1.4, "discussion": 1.4, "experiment": 0.3, "data": 0.3}),
    Question("dataset", "Dữ liệu (dataset)",
             "Bài dùng những bộ dữ liệu nào (tên, quy mô, nguồn, cách chia train/test)? Dữ liệu có công khai không?",
             r"\b(dataset|datasets|data set|benchmark|corpus|database|cohort|subjects|participants|patients|images|"
             r"samples|scans|collected|acquired|split|training set|test set|validation)\b",
             {"data": 2.2, "experiment": 1.6, "method": 1.0}),
    Question("metrics", "Cách đánh giá (thước đo)",
             "Kết quả được đánh giá bằng những thước đo/thí nghiệm nào?",
             r"\b(metric|metrics|evaluat|measure|accuracy|precision|recall|f1|dice|iou|auc|rmse|mae|bleu|psnr|"
             r"ssim|error|distance|correlation|ablation|statistical)\w*",
             {"experiment": 2.0, "method": 1.0}),
    Question("baseline", "So sánh với baseline nào",
             "Bài so sánh với những phương pháp nền (baseline) nào? Kết quả so sánh ra sao?",
             r"\b(baselines?|compar\w*|state[- ]of[- ]the[- ]art|sota|outperform\w*|versus|vs\.?|than|competing|"
             r"existing methods?|previous methods?|methods? such as)\b",
             {"experiment": 2.0, "abstract": 1.2, "related": 1.0}),
    Question("improvement", "Cải thiện được gì",
             "Phương pháp cải thiện được gì so với trước (con số cụ thể: độ chính xác, tốc độ, chi phí...)?",
             r"(\bimprov\w*|\boutperform\w*|\breduc\w*|\bincreas\w*|\bachiev\w*|\bgain\w*|\bfaster\b|\bspeed\w*|"
             r"\d+(\.\d+)?\s?%|\d+(\.\d+)?\s?×|\btimes\b|\bsignificant\w*)",
             {"abstract": 1.8, "experiment": 1.8, "conclusion": 1.5}),
    Question("pros", "Ưu điểm",
             "Ưu điểm nổi bật của phương pháp là gì?",
             r"\b(advantage|benefit|robust|efficient|effective|accurate|fast|lightweight|scalable|generali[sz]|"
             r"flexible|simple|interpretable|superior|reliab)\w*",
             {"abstract": 1.4, "rationale": 1.8, "discussion": 1.6, "conclusion": 1.6, "experiment": 1.2}),
    Question("cons", "Nhược điểm / hạn chế",
             "Phương pháp còn hạn chế/nhược điểm gì (tác giả thừa nhận hoặc thấy rõ từ thí nghiệm)?",
             r"\b(limitation|limitations|limited|fails?|failure|drawback|weakness|cannot|can not|not able|only|"
             r"restricted|requires?|depend|sensitive|expensive|future work|remain|caveat|although)\w*",
             {"discussion": 2.2, "conclusion": 1.6, "rationale": 1.3, "experiment": 1.0}),
    Question("code", "Mã nguồn công khai",
             "Mã nguồn / mô hình / dữ liệu có được công bố công khai không? Ở đâu?",
             r"(https?://|github|gitlab|code (is|will be)|publicly available|open[- ]source|available at|"
             r"released|zenodo|hugging ?face)",
             {"abstract": 1.5, "intro": 1.0, "conclusion": 1.2, "other": 1.0}, k=3),
    Question("future", "Hướng phát triển",
             "Tác giả đề xuất hướng phát triển tiếp theo nào?",
             r"\b(future work|in the future|future|extend|extension|plan to|further|next step|open question)\w*",
             {"discussion": 2.0, "conclusion": 2.0}, k=3),
]


@dataclass
class Evidence:
    text: str  # câu tiếng Anh (công thức dạng LaTeX)
    section: str
    score: float


@dataclass
class Answer:
    question: Question
    text: str  # câu trả lời (Markdown, tiếng Việt)
    evidence: list[Evidence] = field(default_factory=list)


@dataclass
class Summary:
    title: str
    answers: list[Answer]
    facts: dict[str, list[str]]
    mode: str  # "llm" | "extract"
    model: str = ""


_CAPTION = re.compile(r"^\s*(FIGURE|Figure|Fig\.|TABLE|Table)\s*[A-Z]?\d", re.I)


def _paragraphs(s: Section) -> list[str]:
    """Đoạn văn của mục; nối lại đoạn bị cắt ở chỗ sang cột/sang trang (chưa hết câu)."""
    out: list[str] = []
    for r in s.regions:
        if r.kind != "text":
            continue
        t = with_latex(r)
        if out and not re.search(r"[.!?:]\W*$", out[-1]) and t[:1].islower() and not _CAPTION.match(out[-1]):
            out[-1] = out[-1].rstrip("-") + ("" if out[-1].endswith("-") else " ") + t
        else:
            out.append(t)
    return out


def _sentences(sections: list[Section]) -> list[tuple[str, Section, int]]:
    out: list[tuple[str, Section, int]] = []
    for s in sections:
        for para in _paragraphs(s):
            for i, sent in enumerate(split_sentences(para, max_len=600)):
                if len(sent) >= 25 and re.search(r"[a-z]{3}", sent):
                    out.append((sent, s, i))
    return out


def select_evidence(q: Question, sents: list[tuple[str, Section, int]], k: int) -> list[Evidence]:
    pat = re.compile(q.pattern, re.I)
    scored: list[Evidence] = []
    for sent, sec, idx in sents:
        hits = len(pat.findall(sent))
        if not hits:
            continue
        w = q.sections.get(sec.kind, 0.5)
        if _CAPTION.match(sent):
            w *= 0.6 if q.key in ("architecture", "io", "baseline") else 0.3
        score = w * (1 + min(hits, 4)) * (1.1 if idx == 0 else 1.0) / (1 + len(sent) / 600)
        scored.append(Evidence(sent, sec.title, score))
    scored.sort(key=lambda e: -e.score)
    seen: set[str] = set()
    out = []
    for e in scored:
        key = e.text[:60]
        if key not in seen:
            seen.add(key)
            out.append(e)
        if len(out) >= k:
            break
    return out


# ---------------------------------------------------------------------------
# Dữ kiện tự nhận ra (không cần LLM)
# ---------------------------------------------------------------------------

_URL_RE = re.compile(r"https?://[^\s)\]>,;]+|(?:www\.)?github\.com/[^\s)\]>,;]+", re.I)
_CODE_HOSTS = re.compile(r"github|gitlab|bitbucket|huggingface|zenodo|codeocean|osf\.io|figshare|kaggle|sourceforge", re.I)
_DATASET_RE = re.compile(
    r"\b((?:[A-Z][A-Za-z0-9]*(?:[-_][A-Za-z0-9]+)*)(?:\s+[A-Z0-9][A-Za-z0-9-]*){0,3})\s+"
    r"(?:dataset|data set|benchmark|database|cohort|corpus|challenge)\b"
    r"|\b(?:dataset|benchmark|database|cohort)s?\s+(?:called|named|termed)?\s*\(?((?:[A-Z][A-Za-z0-9]*[-_]?)+[A-Za-z0-9]*)")
_NOT_NAMES = {"The", "This", "Our", "We", "Each", "All", "These", "Both", "Public", "Large", "Real", "Test", "Training",
              "Validation", "Clinical", "In", "For", "A", "An", "Two", "Three", "Four", "Five", "Several", "Other", "New",
              "Table", "Figure", "Fig", "Section", "Synthetic", "Additional", "Independent", "External", "Internal",
              "Original", "Full", "Whole", "Entire", "Same", "Different", "Second", "First", "Third", "Final"}
_METHOD_TOKEN = re.compile(r"(?<![\w-])([A-Z][A-Za-z]*[A-Z0-9][A-Za-z0-9-]*|[A-Z]{2,}[A-Za-z0-9-]*|"
                           r"[a-z]+[A-Z][A-Za-z0-9-]*)(?![\w])")


# Thước đo, phần cứng, cặp ngôn ngữ, đơn vị: hay đứng cạnh tên phương pháp trong câu so sánh nhưng không phải.
_NOT_METHOD = re.compile(
    r"^(BLEU|ROUGE(-\w+)?|METEOR|CIDE[Rr]|PSNR|SSIM|LPIPS|FID|FLOPS?|FLOPs|GFLOPs|MACs|AUC|ROC|AUROC|MSE|RMSE|MAE|"
    r"MAPE|NLL|PPL|WER|CER|mAP|mIoU|IoU|HD\d*|ASSD|Dice|DSC|F1|SOTA|GPUs?|TPUs?|CPUs?|RAM|VRAM|"
    r"[KPVATH]\d{2,4}|RTX\w*|GTX\w*|NVIDIA|[A-Z]{2}-[A-Z]{2}|WMT\d*|ImageNet\w*|[A-Z][a-z]?\d?)$")


def _method_like(tok: str) -> bool:
    """Tên phương pháp: CamelCase / có số / gạch nối (CoSeg, T1w-Reg, ResNet-50) hoặc viết hoa ≥ 4 chữ."""
    if tok in _NOT_NAMES or tok.upper() in ("FIGURE", "TABLE", "SECTION", "EQ") or tok.isdigit() or len(tok) > 25:
        return False
    if _NOT_METHOD.match(tok):
        return False
    if tok.isupper():
        return len(tok) >= 4  # bỏ FA, MD, WM, CSF... (thường là đại lượng, không phải phương pháp)
    return True


# Link bị xuống dòng trong PDF: "https://github.com/ tensorflow/tensor2tensor" → nối lại.
_URL_BREAK = re.compile(r"((?:https?://|www\.|github\.com/)[^\s]*[/_-])\s+(?=[a-z0-9][\w.-]*(?:/|[\w-]*[_-][\w-]))")


def _join_broken_urls(text: str) -> str:
    return _URL_BREAK.sub(r"\1", text)


def extract_facts(ir: Document, sections: list[Section]) -> dict[str, list[str]]:
    all_text = _join_broken_urls(" ".join(r.plain for r in ir.regions))
    urls = []
    for u in _URL_RE.findall(all_text):
        u = u.rstrip(".")
        if not re.search(r"\.[a-z]{2,}/[\w-]", u, re.I):
            continue  # chỉ có tên miền ("https://github.com/") → không phải link mã nguồn cụ thể
        if _CODE_HOSTS.search(u) and u not in urls:
            urls.append(u)
    body = " ".join(r.plain for s in sections for r in s.regions if r.kind == "text")
    names: Counter[str] = Counter()
    for m in _DATASET_RE.finditer(body):
        name = (m.group(1) or m.group(2) or "").strip()
        words = name.split()
        while words and words[0] in _NOT_NAMES:
            words = words[1:]
        name = " ".join(words)
        if len(name) >= 3 and not name.isdigit() and name not in _NOT_NAMES:
            names[name] += 1
    # Phương pháp so sánh: tên riêng dạng viết hoa/CamelCase trong các câu so sánh và trong bảng
    methods: Counter[str] = Counter()
    for sent in split_sentences(body, max_len=800):
        if re.search(r"\b(compar\w*|baseline\w*|outperform\w*|state[- ]of[- ]the[- ]art|versus|against)\b", sent, re.I):
            for tok in _METHOD_TOKEN.findall(sent):
                if _method_like(tok):
                    methods[tok] += 1
    for r in ir.regions:
        if r.kind == "table":
            for line in r.plain.splitlines():
                first = line.strip().split("  ")[0].split(" ")[0] if line.strip() else ""
                if _METHOD_TOKEN.fullmatch(first or "-") and _method_like(first):
                    methods[first] += 1
    dataset_names = [n for n, _ in names.most_common(8)]
    return {
        "code": urls[:6],
        "datasets": dataset_names,
        "baselines": [m for m, c in methods.most_common(12) if m not in dataset_names][:10],
    }


# ---------------------------------------------------------------------------
# Tóm tắt
# ---------------------------------------------------------------------------

SUMMARY_SYSTEM = """Bạn là nhà khoa học đọc và tóm tắt bài báo nghiên cứu cho đồng nghiệp người Việt.
Quy tắc:
- Chỉ dùng thông tin trong TRÍCH ĐOẠN được cung cấp; không bịa.
- Trích đoạn chỉ trả lời được một phần → vẫn trả lời phần đó. Chỉ khi trích đoạn hoàn toàn không liên quan đến câu hỏi mới viết đúng một dòng: "Bài báo không nêu rõ."
- Trả lời bằng tiếng Việt, súc tích, 2–6 gạch đầu dòng ("- "), có con số cụ thể nếu có.
- Giữ nguyên tên riêng, tên phương pháp, tên dataset, đường link; công thức viết bằng LaTeX $...$.
- Không lặp lại câu hỏi, không mở đầu bằng lời dẫn."""


RETRY_SYSTEM = SUMMARY_SYSTEM.split("\n- Trích đoạn chỉ")[0] + """
- Các trích đoạn đã được chọn vì có liên quan đến câu hỏi: hãy tổng hợp câu trả lời từ chúng.
- Trả lời bằng tiếng Việt, súc tích, 2–6 gạch đầu dòng ("- "), có con số cụ thể nếu có.
- Giữ nguyên tên riêng, tên phương pháp, tên dataset, đường link; công thức viết bằng LaTeX $...$."""

TRANSLATE_SYSTEM = ("Dịch các câu tiếng Anh sau sang tiếng Việt, mỗi câu một dòng bắt đầu bằng \"- \". Giữ nguyên tên "
                    "riêng, số liệu, đường link, công thức $...$. Chỉ trả về các dòng đã dịch.")


def _evidence_fallback(llm, ev: list[Evidence]) -> str:
    note = "_Mô hình không tự tổng hợp được câu trả lời; các câu liên quan nhất trong bài:_\n"
    try:
        vi = _clean_answer(llm._chat(TRANSLATE_SYSTEM, "\n".join(f"- {e.text}" for e in ev)))
    except Exception:  # noqa: BLE001
        vi = ""
    if not vi or _is_refusal(vi):
        vi = "\n".join(f"- {e.text}" for e in ev)
    return note + vi


def _clean_answer(text: str) -> str:
    """Bỏ thẻ mục bị chép từ trích đoạn ("- [1 Introduction] ...") và chữ Hán lẫn vào (qwen)."""
    text = re.sub(r"(?m)^(\s*[-*]\s*)\[[^\]\n]{1,60}\]\s*", r"\1", text or "")
    text = text.replace("表", "Bảng ").replace("图", "Hình ")
    text = re.sub(r"[\u3000-\u303f\u4e00-\u9fff\uff00-\uffef]+", "", text)
    return re.sub(r"[ \t]{2,}", " ", text).strip()


RETRY_HINT = ("\n\nLưu ý: các trích đoạn trên được chọn vì có liên quan đến câu hỏi. Hãy trả lời dựa trên chúng, "
              "kể cả khi chỉ trả lời được một phần (nêu rõ phần nào bài có nói).")


def _is_refusal(text: str) -> bool:
    t = re.sub(r"[\s\-*_.\"“”]+", " ", text).strip().lower()
    return not t or (len(t) < 80 and "không nêu rõ" in t)


def _llm_context(q: Question, ev: list[Evidence], abstract: str, facts: dict[str, list[str]]) -> str:
    parts = []
    if abstract:
        parts.append("TÓM TẮT CỦA BÀI (Abstract):\n" + abstract[:2500])
    parts.append("TRÍCH ĐOẠN LIÊN QUAN:\n" + "\n".join(f"- [{e.section}] {e.text}" for e in ev))
    if q.key == "code" and facts.get("code"):
        parts.append("ĐƯỜNG LINK TÌM THẤY TRONG BÀI: " + ", ".join(facts["code"]))
    if q.key == "dataset" and facts.get("datasets"):
        parts.append("TÊN DATASET TỰ NHẬN RA (có thể sai): " + ", ".join(facts["datasets"]))
    if q.key == "baseline" and facts.get("baselines"):
        parts.append("TÊN PHƯƠNG PHÁP XUẤT HIỆN KHI SO SÁNH (có thể sai): " + ", ".join(facts["baselines"]))
    return "\n\n".join(parts) + f"\n\nCÂU HỎI: {q.ask}"


def summarize(
    ir: Document,
    llm=None,
    translate: Callable[[list[str]], list[str]] | None = None,
    status: StatusFn | None = None,
) -> Summary:
    """Tóm tắt bài theo bộ câu hỏi. ``llm``: LLMTranslator đang chạy (hoặc None);
    ``translate``: hàm dịch Anh→Việt cho chế độ trích câu."""
    sections = split_sections(ir)
    sents = _sentences(sections)
    facts = extract_facts(ir, sections)
    abstract = " ".join(with_latex(r) for s in sections if s.kind == "abstract" for r in s.regions if r.kind == "text")
    title = _guess_title(ir)
    answers: list[Answer] = []
    for i, q in enumerate(QUESTIONS, 1):
        if status:
            status(f"Tóm tắt: {q.title} ({i}/{len(QUESTIONS)})")
        if llm is not None:
            ev = select_evidence(q, sents, 14)
            try:
                prompt = _llm_context(q, ev, abstract, facts)
                text = _clean_answer(llm._chat(SUMMARY_SYSTEM, prompt)) if (ev or abstract) else ""
                if _is_refusal(text) and len(ev) >= 3:
                    # Mô hình nhỏ hay né trả lời dù trích đoạn có thông tin → hỏi lại một lần, không cho
                    # lối "không nêu rõ"; vẫn né thì đưa các câu liên quan nhất (nhờ LLM dịch).
                    text = _clean_answer(llm._chat(RETRY_SYSTEM, prompt + RETRY_HINT)) or text
                    if _is_refusal(text):
                        text = _evidence_fallback(llm, ev[:3])
            except Exception as e:  # noqa: BLE001  (LLM lỗi giữa chừng → trích câu cho câu hỏi này)
                log.warning("LLM lỗi khi tóm tắt '%s': %s", q.title, e)
                text = ""
            answers.append(Answer(q, text.strip() or "Bài báo không nêu rõ.", ev[:3]))
        else:
            answers.append(Answer(q, "", select_evidence(q, sents, q.k)))
    if llm is None and translate is not None:
        flat = [e.text for a in answers for e in a.evidence]
        if flat:
            if status:
                status(f"Tóm tắt: đang dịch {len(flat)} câu trích dẫn...")
            vi = translate(flat)
            it = iter(vi)
            for a in answers:
                a.text = "\n".join(f"- {next(it)}" for _ in a.evidence) or "_Không tìm thấy câu liên quan trong bài._"
    elif llm is None:
        for a in answers:
            a.text = "\n".join(f"- {e.text}" for e in a.evidence) or "_Không tìm thấy câu liên quan trong bài._"
    _add_fact_lines(answers, facts)
    from .config import settings

    return Summary(title, answers, facts, "llm" if llm is not None else "extract",
                   settings.llm_model if llm is not None else "")


def _add_fact_lines(answers: list[Answer], facts: dict[str, list[str]]) -> None:
    for a in answers:
        extra = ""
        if a.question.key == "code":
            extra = ("**Link tìm thấy trong bài:** " + ", ".join(f"<{u}>" for u in facts["code"])) if facts["code"] \
                else "**Không tìm thấy link mã nguồn (GitHub, GitLab, Zenodo…) trong bài.**"
        elif a.question.key == "dataset" and facts["datasets"]:
            extra = "**Tên dataset nhận ra:** " + ", ".join(facts["datasets"])
        elif a.question.key == "baseline" and facts["baselines"]:
            extra = "**Phương pháp xuất hiện trong phần so sánh:** " + ", ".join(facts["baselines"])
        if extra:
            a.text = (a.text + "\n\n" + extra).strip()


def _guess_title(ir: Document) -> str:
    if not ir.pages:
        return ir.source
    first = [r for r in ir.pages[0].regions if r.kind == "text" and 10 < len(r.plain) < 250]
    if not first:
        return ir.source
    r = max(first, key=lambda r: (r.font_size, -r.bbox.y0))
    return re.sub(r"\s+", " ", r.plain).strip()


# ---------------------------------------------------------------------------
# Xuất kết quả
# ---------------------------------------------------------------------------


def summary_markdown(s: Summary) -> str:
    how = (f"Tạo bằng LLM cục bộ `{s.model}` — chỉ dựa trên nội dung bài; nên đối chiếu với bản gốc."
           if s.mode == "llm" else
           "Chế độ trích câu (chưa có LLM cục bộ): mỗi mục là các câu liên quan nhất trích từ bài và dịch sang tiếng "
           "Việt. Bật Ollama/LM Studio để có câu trả lời tổng hợp.")
    out = [f"# Tóm tắt bài báo\n\n**{s.title}**\n\n_{how}_\n"]
    for a in s.answers:
        out.append(f"## {a.question.title}\n\n{a.text}\n")
        if s.mode == "llm" and a.evidence:
            out.append("<details><summary>Trích đoạn gốc</summary>\n\n" +
                       "\n".join(f"- [{e.section}] {e.text}" for e in a.evidence) + "\n\n</details>\n")
    return "\n".join(out).strip() + "\n"


def markdown_to_html(md: str, title: str) -> str:
    """Markdown đơn giản (tiêu đề, gạch đầu dòng, đậm, link, công thức) → trang HTML xem trên trình duyệt."""
    body: list[str] = []
    in_list = False
    for line in md.splitlines():
        raw = line.rstrip()
        if raw.startswith("<details>") or raw.startswith("</details>"):
            if in_list:
                body.append("</ul>")
                in_list = False
            body.append(raw.replace("<summary>", "<summary>").replace("</summary>", "</summary>"))
            continue
        esc = html.escape(raw, quote=False)
        esc = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", esc)
        esc = re.sub(r"(?<![\w*])_(.+?)_(?![\w*])", r"<i>\1</i>", esc)
        esc = re.sub(r"&lt;(https?://[^&\s]+)&gt;", r'<a href="\1" target="_blank">\1</a>', esc)
        esc = re.sub(r"`([^`]+)`", r"<code>\1</code>", esc)
        if raw.startswith("- "):
            if not in_list:
                body.append("<ul>")
                in_list = True
            body.append(f"<li>{esc[2:]}</li>")
            continue
        if in_list:
            body.append("</ul>")
            in_list = False
        if raw.startswith("### "):
            body.append(f"<h3>{esc[4:]}</h3>")
        elif raw.startswith("## "):
            body.append(f"<h2>{esc[3:]}</h2>")
        elif raw.startswith("# "):
            body.append(f"<h1>{esc[2:]}</h1>")
        elif raw.startswith("$$") and raw.endswith("$$") and len(raw) > 4:
            body.append(f"<p class='eq'>{esc}</p>")
        elif raw:
            body.append(f"<p>{esc}</p>")
    if in_list:
        body.append("</ul>")
    return f"""<!doctype html><html lang="vi"><head><meta charset="utf-8"><title>{html.escape(title)}</title>
<meta name="viewport" content="width=device-width, initial-scale=1">
<script>window.MathJax={{tex:{{inlineMath:[['$','$']]}}}};</script>
<script async src="https://cdn.jsdelivr.net/npm/mathjax@3.2.2/es5/tex-chtml.js"></script>
<style>body{{font-family:Georgia,'Times New Roman',serif;max-width:860px;margin:24px auto;padding:0 16px;line-height:1.6;color:#1f2933}}
h1{{font-size:1.6em}} h2{{font-size:1.2em;margin-top:1.6em;border-bottom:1px solid #e3e8ee;padding-bottom:4px}}
li{{margin:4px 0}} details{{color:#52606d;font-size:.92em;margin:6px 0 0}} .eq{{text-align:center}}
code{{background:#f1f4f8;padding:1px 4px;border-radius:3px}}</style></head><body>
{chr(10).join(body)}
</body></html>"""


def write_docx(md: str, path, title: str) -> None:
    """Markdown đơn giản → Word."""
    from docx import Document as Docx
    from docx.shared import Pt

    doc = Docx()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)
    skip = False
    for line in md.splitlines():
        raw = line.rstrip()
        if raw.startswith("<details>"):
            skip = True
            continue
        if raw.startswith("</details>"):
            skip = False
            continue
        if skip or not raw:
            continue
        if raw.startswith("# "):
            doc.add_heading(raw[2:], level=0)
        elif raw.startswith("## "):
            doc.add_heading(raw[3:], level=1)
        elif raw.startswith("### "):
            doc.add_heading(raw[4:], level=2)
        else:
            style = "List Bullet" if raw.startswith("- ") else None
            par = doc.add_paragraph(style=style)
            text = raw[2:] if raw.startswith("- ") else raw
            for k, piece in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
                piece = re.sub(r"<(https?://[^>]+)>", r"\1", piece)
                if piece.startswith("_") and piece.endswith("_") and len(piece) > 2:
                    run = par.add_run(piece[1:-1])
                    run.italic = True
                else:
                    run = par.add_run(piece)
                run.bold = k % 2 == 1 or None
    doc.core_properties.title = title
    doc.save(str(path))

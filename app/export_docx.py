"""Xuất bản dịch ra Word (.docx) để chỉnh sửa tiếp.

- Thứ tự đọc: phần rộng hết trang (tiêu đề, tóm tắt) theo chiều dọc; phần 2 cột đọc hết
  cột trái rồi tới cột phải.
- Tiêu đề / đoạn văn / tài liệu tham khảo là chữ (sửa được); chỉ số trên, công thức nội
  dòng in nghiêng, đoạn nghi ngờ tô vàng.
- Công thức đứng riêng, bảng, hình được cắt từ PDF gốc thành ảnh độ phân giải cao.
"""

from __future__ import annotations

import io
import re
from pathlib import Path

import pymupdf

from .ir import Document, Region
from .pipeline import MATH_END, MATH_START, SUP_END, SUP_START
from .scientific import translate_heading

_AS_IMAGE = {"math", "table", "figure"}
_SKIP = {"rotated", "abandon", "empty", "data", "code"}
_NUMBERED = re.compile(r"^\s*(\d+(?:\.\d+)*)[.\s|]")


def reading_order(items: list[tuple[pymupdf.Rect, object]], page_width: float) -> list:
    """Sắp xếp theo thứ tự đọc: dải rộng hết trang ngắt mạch; trong mỗi đoạn 2 cột đọc trái rồi phải."""
    mid = page_width / 2
    out: list = []
    group: list[tuple[pymupdf.Rect, object]] = []

    def flush() -> None:
        left = sorted((x for x in group if (x[0].x0 + x[0].x1) / 2 < mid), key=lambda x: x[0].y0)
        right = sorted((x for x in group if (x[0].x0 + x[0].x1) / 2 >= mid), key=lambda x: x[0].y0)
        out.extend(obj for _, obj in left + right)
        group.clear()

    for rect, obj in sorted(items, key=lambda x: (x[0].y0, x[0].x0)):
        if rect.width > 0.6 * page_width or (rect.x0 < mid - 20 and rect.x1 > mid + 20):
            flush()
            out.append(obj)
        else:
            group.append((rect, obj))
    flush()
    return out


def _add_runs(par, text: str, math_spans: list[str], highlight: bool, base_size: float | None = None) -> None:
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import Pt

    text = text.replace(MATH_START, "").replace(MATH_END, "")
    # Đánh dấu vùng in nghiêng (công thức nội dòng) bằng chỉ mục ký tự
    italic = [False] * len(text)
    for m in sorted(set(math_spans), key=len, reverse=True):
        if len(m) == 1 and not m.isalpha():
            continue
        pat = re.escape(m)
        if m[:1].isalnum():
            pat = r"(?<![\w])" + pat
        if m[-1:].isalnum():
            pat += r"(?![\w])"
        for hit in re.finditer(pat, text):
            for i in range(hit.start(), hit.end()):
                italic[i] = True
    # Tách theo chỉ số trên và kiểu nghiêng
    chunks: list[tuple[str, bool, bool]] = []
    sup = False
    buf, buf_it = "", None
    for ch, it in zip(text, italic):
        if ch in (SUP_START, SUP_END):
            if buf:
                chunks.append((buf, bool(buf_it), sup))
            buf, buf_it = "", None
            sup = ch == SUP_START
            continue
        if buf and it != buf_it:
            chunks.append((buf, bool(buf_it), sup))
            buf = ""
        buf += ch
        buf_it = it
    if buf:
        chunks.append((buf, bool(buf_it), sup))
    for chunk, it, is_sup in chunks:
        run = par.add_run(chunk)
        run.italic = it or None
        run.font.superscript = is_sup or None
        if base_size:
            run.font.size = Pt(base_size)
        if highlight:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _crop_png(page: pymupdf.Page, rect: pymupdf.Rect, dpi: int = 200) -> bytes:
    clip = pymupdf.Rect(rect) + (-2, -2, 2, 2)
    return page.get_pixmap(clip=clip & page.rect, dpi=dpi).tobytes("png")


def _heading_level(r: Region, body_size: float) -> int | None:
    text = r.plain.strip()
    if len(text) > 160 or len(r.line_rects) > 3:
        return None
    if translate_heading(text) is not None:
        m = _NUMBERED.match(text)
        return min(3, m.group(1).count(".") + 1) if m else 1
    if r.font_size >= body_size * 1.5:
        return 0
    m = _NUMBERED.match(text)
    if m and (r.bold or r.font_size >= body_size * 1.1) and len(text) < 120 and not text.endswith("."):
        return min(3, m.group(1).count(".") + 1)  # mục đánh số: "2.1 Pseudo-ground-truth"
    if r.font_size >= body_size * 1.15 and len(text) < 120 and not text.endswith("."):
        return 2
    return None


def build_docx(ir: Document, source_pdf: pymupdf.Document, out_path: Path, highlight: bool = True,
               is_doubtful=lambda r: False) -> Path:
    from docx import Document as Docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, Inches

    doc = Docx()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    section = doc.sections[0]
    usable = section.page_width - section.left_margin - section.right_margin

    sizes = sorted(r.font_size for r in ir.regions if r.kind == "text" and len(r.plain) > 200)
    body_size = sizes[len(sizes) // 2] if sizes else 10.0

    for p in ir.pages:
        page = source_pdf[p.index]
        items: list[tuple[pymupdf.Rect, object]] = []
        for r in p.regions:
            if r.kind in _SKIP:
                continue
            items.append((pymupdf.Rect(r.bbox), r))
        # Ảnh nhúng (hình vẽ dạng ảnh) chưa nằm trong vùng nào
        regions_rects = [pymupdf.Rect(r.bbox) for r in p.regions if r.kind in _AS_IMAGE]
        for info in page.get_image_info():
            rect = pymupdf.Rect(info["bbox"]) & page.rect
            if p.scanned or rect.width < 40 or rect.height < 40:
                continue
            if any(abs(rect & rr) > 0.5 * abs(rect) for rr in regions_rects):
                continue
            items.append((rect, ("image", rect)))

        for obj in reading_order(items, page.rect.width):
            if isinstance(obj, tuple):  # ảnh nhúng
                rect = obj[1]
                width = min(usable, int(usable * rect.width / page.rect.width * 1.4))
                doc.add_picture(io.BytesIO(_crop_png(page, rect)), width=width)
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                continue
            r: Region = obj
            if r.kind in _AS_IMAGE:
                width = min(usable, int(usable * r.bbox.width / page.rect.width * 1.4))
                doc.add_picture(io.BytesIO(_crop_png(page, r.bbox)), width=max(width, Inches(0.5)))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                continue
            text = r.translation if r.translation is not None else r.text
            if r.kind == "reference":
                par = doc.add_paragraph()
                _add_runs(par, text, [], False, base_size=9)
                continue
            level = _heading_level(r, body_size)
            if level is not None:
                h = doc.add_heading(level=level)
                _add_runs(h, text, r.math_spans, highlight and is_doubtful(r))
                continue
            par = doc.add_paragraph()
            if r.align == "center":
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif r.align == "justify":
                par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            small = 9 if r.font_size < body_size * 0.9 else None
            _add_runs(par, text, r.math_spans, highlight and is_doubtful(r), base_size=small)
            if r.bold:
                for run in par.runs:
                    run.bold = True

    doc.core_properties.title = Path(ir.source).stem
    doc.core_properties.comments = "Dịch bằng PDF Translate (chạy cục bộ)"
    doc.save(str(out_path))
    return out_path
